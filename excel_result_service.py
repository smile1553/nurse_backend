import hashlib
import json
import os
import re
import tempfile
import threading
import time
import uuid
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterator, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


RESULT_SHEET = "學生結果"
REQUEST_SHEET = "_requests"
RESULT_HEADERS = ["登入時間", "學號", "答對題數", "題目分數", "語氣分數", "總分"]
REQUEST_HEADERS = [
    "resultId",
    "payloadHash",
    "rowNumber",
    "submittedAt",
    "studentId",
    "loginTime",
    "correctCount",
    "toneScore",
    "questionScore",
    "totalScore",
]
SESSION_ID_RE = re.compile(r"^[A-Za-z0-9_-]{1,80}$")


class ResultServiceError(Exception):
    pass


class SessionNotFoundError(ResultServiceError):
    pass


class InactiveSessionError(ResultServiceError):
    pass


class IdempotencyConflictError(ResultServiceError):
    pass


class StorageBusyError(ResultServiceError):
    pass


class ExcelResultService:
    def __init__(self, output_dir: Path, lock_timeout_seconds: float = 10.0):
        self.output_dir = Path(output_dir).resolve()
        self.sessions_dir = self.output_dir / "sessions"
        self.locks_dir = self.output_dir / "locks"
        self.current_session_path = self.output_dir / "current_session.json"
        self.lock_timeout_seconds = lock_timeout_seconds
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.sessions_dir.mkdir(parents=True, exist_ok=True)
        self.locks_dir.mkdir(parents=True, exist_ok=True)
        self._global_lock = threading.RLock()
        self._session_locks: Dict[str, threading.RLock] = {}

    def create_session(self, now: Optional[datetime] = None) -> Dict[str, str]:
        created_at = now or datetime.now().astimezone()
        if created_at.tzinfo is None or created_at.utcoffset() is None:
            created_at = created_at.astimezone()

        with self._global_lock, self._file_lock(self.locks_dir / "active_session.lock"):
            while True:
                session_id = f"{created_at:%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"
                workbook_path = self.output_dir / f"session_{session_id}.xlsx"
                if not workbook_path.exists():
                    break

            workbook = self._new_workbook()
            self._save_workbook_atomic(workbook, workbook_path)

            metadata = {
                "sessionId": session_id,
                "createdAt": created_at.isoformat(timespec="seconds"),
                "workbook": workbook_path.name,
            }
            self._write_json_atomic(self.sessions_dir / f"{session_id}.json", metadata)
            self._write_json_atomic(self.current_session_path, metadata)
            return dict(metadata)

    def get_current_session(self) -> Optional[Dict[str, str]]:
        if not self.current_session_path.exists():
            return None
        try:
            data = json.loads(self.current_session_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ResultServiceError("current session metadata is unreadable") from exc
        self._validate_metadata(data)
        return data

    def ensure_active_session(self, session_id: str) -> Dict[str, str]:
        self._validate_session_id(session_id)
        metadata_path = self.sessions_dir / f"{session_id}.json"
        if not metadata_path.exists():
            raise SessionNotFoundError("session not found")

        current = self.get_current_session()
        if current is None or current.get("sessionId") != session_id:
            raise InactiveSessionError("session is no longer active")

        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ResultServiceError("session metadata is unreadable") from exc
        self._validate_metadata(metadata)
        workbook_path = self._workbook_path(metadata)
        if not workbook_path.exists():
            raise SessionNotFoundError("session workbook not found")
        return metadata

    def append_result(
        self,
        session_id: str,
        result_id: str,
        student_id: str,
        login_time: str,
        correct_count: int,
        tone_score: int,
    ) -> Dict[str, Any]:
        normalized_result_id = str(uuid.UUID(str(result_id)))
        payload = {
            "studentId": student_id,
            "loginTime": login_time,
            "correctCount": correct_count,
            "toneScore": tone_score,
        }
        payload_hash = hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest()
        question_score = correct_count * 10
        total_score = question_score + tone_score

        thread_lock = self._session_lock(session_id)
        active_lock_path = self.locks_dir / "active_session.lock"
        file_lock_path = self.locks_dir / f"{session_id}.lock"
        with (
            self._global_lock,
            self._file_lock(active_lock_path),
            thread_lock,
            self._file_lock(file_lock_path),
        ):
            metadata = self.ensure_active_session(session_id)
            workbook_path = self._workbook_path(metadata)
            try:
                workbook = load_workbook(workbook_path)
            except Exception as exc:
                raise ResultServiceError("failed to open session workbook") from exc

            request_sheet = workbook[REQUEST_SHEET]
            duplicate = self._find_request(request_sheet, normalized_result_id)
            if duplicate is not None:
                if duplicate["payloadHash"] != payload_hash:
                    raise IdempotencyConflictError("resultId was already used with different data")
                return {
                    "ok": True,
                    "duplicate": True,
                    "sessionId": session_id,
                    "resultId": normalized_result_id,
                    "correctCount": duplicate["correctCount"],
                    "questionScore": duplicate["questionScore"],
                    "toneScore": duplicate["toneScore"],
                    "totalScore": duplicate["totalScore"],
                }

            result_sheet = workbook[RESULT_SHEET]
            result_sheet.append(
                [
                    login_time,
                    student_id,
                    f"{correct_count}/8",
                    f"{question_score}/80",
                    f"{tone_score}/20",
                    f"{total_score}/100",
                ]
            )
            row_number = result_sheet.max_row
            self._format_result_row(result_sheet, row_number)
            result_sheet.auto_filter.ref = f"A1:F{row_number}"

            submitted_at = datetime.now().astimezone().isoformat(timespec="seconds")
            request_sheet.append(
                [
                    normalized_result_id,
                    payload_hash,
                    row_number,
                    submitted_at,
                    student_id,
                    login_time,
                    correct_count,
                    tone_score,
                    question_score,
                    total_score,
                ]
            )

            self._save_workbook_atomic(workbook, workbook_path)
            return {
                "ok": True,
                "duplicate": False,
                "sessionId": session_id,
                "resultId": normalized_result_id,
                "correctCount": correct_count,
                "questionScore": question_score,
                "toneScore": tone_score,
                "totalScore": total_score,
            }

    def _new_workbook(self) -> Workbook:
        workbook = Workbook()
        result_sheet = workbook.active
        result_sheet.title = RESULT_SHEET
        result_sheet.append(RESULT_HEADERS)
        result_sheet.freeze_panes = "A2"
        result_sheet.auto_filter.ref = "A1:F1"
        result_sheet.sheet_view.showGridLines = False

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        thin_white = Side(style="thin", color="FFFFFF")
        for cell in result_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(right=thin_white)
        result_sheet.row_dimensions[1].height = 24

        widths = [28, 18, 14, 14, 14, 14]
        for index, width in enumerate(widths, start=1):
            result_sheet.column_dimensions[get_column_letter(index)].width = width

        request_sheet = workbook.create_sheet(REQUEST_SHEET)
        request_sheet.append(REQUEST_HEADERS)
        request_sheet.sheet_state = "hidden"
        return workbook

    @staticmethod
    def _format_result_row(sheet: Any, row_number: int) -> None:
        for column in range(1, 7):
            cell = sheet.cell(row=row_number, column=column)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                horizontal="left" if column in (1, 2) else "center",
                vertical="center",
            )
            cell.number_format = "@"
            cell.data_type = "s"
        sheet.row_dimensions[row_number].height = 22

    @staticmethod
    def _find_request(sheet: Any, result_id: str) -> Optional[Dict[str, Any]]:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if str(row[0] or "") != result_id:
                continue
            return {
                "payloadHash": str(row[1]),
                "correctCount": int(row[6]),
                "toneScore": int(row[7]),
                "questionScore": int(row[8]),
                "totalScore": int(row[9]),
            }
        return None

    def _session_lock(self, session_id: str) -> threading.RLock:
        self._validate_session_id(session_id)
        with self._global_lock:
            return self._session_locks.setdefault(session_id, threading.RLock())

    def _workbook_path(self, metadata: Dict[str, str]) -> Path:
        workbook_name = Path(metadata["workbook"]).name
        path = (self.output_dir / workbook_name).resolve()
        if path.parent != self.output_dir:
            raise ResultServiceError("invalid workbook path in session metadata")
        return path

    @staticmethod
    def _validate_session_id(session_id: str) -> None:
        if not SESSION_ID_RE.fullmatch(session_id or ""):
            raise SessionNotFoundError("invalid session id")

    @staticmethod
    def _validate_metadata(metadata: Any) -> None:
        if not isinstance(metadata, dict):
            raise ResultServiceError("invalid session metadata")
        if not SESSION_ID_RE.fullmatch(str(metadata.get("sessionId", ""))):
            raise ResultServiceError("invalid session metadata")
        if not str(metadata.get("workbook", "")).endswith(".xlsx"):
            raise ResultServiceError("invalid session metadata")

    @contextmanager
    def _file_lock(self, lock_path: Path) -> Iterator[None]:
        deadline = time.monotonic() + self.lock_timeout_seconds
        handle: Optional[int] = None
        while handle is None:
            try:
                handle = os.open(str(lock_path), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.write(handle, f"{os.getpid()}\n".encode("ascii"))
            except FileExistsError:
                try:
                    age = time.time() - lock_path.stat().st_mtime
                    if age > 120:
                        lock_path.unlink(missing_ok=True)
                        continue
                except FileNotFoundError:
                    continue
                if time.monotonic() >= deadline:
                    raise StorageBusyError("result workbook is busy; retry the request")
                time.sleep(0.05)
        try:
            yield
        finally:
            if handle is not None:
                os.close(handle)
            lock_path.unlink(missing_ok=True)

    @staticmethod
    def _save_workbook_atomic(workbook: Workbook, target: Path) -> None:
        temporary_path: Optional[Path] = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{target.stem}_",
                suffix=".xlsx",
                dir=target.parent,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, target)
        except PermissionError as exc:
            raise StorageBusyError("Excel file is open or cannot be replaced; retry later") from exc
        except Exception as exc:
            raise ResultServiceError("failed to save session workbook") from exc
        finally:
            if temporary_path is not None and temporary_path.exists():
                temporary_path.unlink(missing_ok=True)

    @staticmethod
    def _write_json_atomic(target: Path, data: Dict[str, Any]) -> None:
        temporary_path: Optional[Path] = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{target.stem}_",
                suffix=".json",
                dir=target.parent,
                delete=False,
                mode="w",
                encoding="utf-8",
            ) as temporary:
                json.dump(data, temporary, ensure_ascii=False, indent=2)
                temporary.flush()
                os.fsync(temporary.fileno())
                temporary_path = Path(temporary.name)
            os.replace(temporary_path, target)
        except Exception as exc:
            raise ResultServiceError("failed to save session metadata") from exc
        finally:
            if temporary_path is not None and temporary_path.exists():
                temporary_path.unlink(missing_ok=True)
