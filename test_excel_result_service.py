import tempfile
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from excel_result_service import (
    ExcelResultService,
    IdempotencyConflictError,
    InactiveSessionError,
    REQUEST_SHEET,
    RESULT_HEADERS,
    RESULT_SHEET,
)
from result_models import StudentResultSubmission


class ExcelResultServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output_dir = Path(self.temp_dir.name) / "student_results"
        self.service = ExcelResultService(self.output_dir, lock_timeout_seconds=2.0)
        self.now = datetime(2026, 9, 21, 10, 30, tzinfo=timezone.utc)
        self.session = self.service.create_session(self.now)

    def tearDown(self):
        self.temp_dir.cleanup()

    def workbook_path(self, session=None):
        session = session or self.session
        return self.output_dir / session["workbook"]

    def append(self, result_id=None, student_id="412345678", correct_count=7, tone_score=18):
        return self.service.append_result(
            self.session["sessionId"],
            str(result_id or uuid.uuid4()),
            student_id,
            "2026-09-21T18:30:00+08:00",
            correct_count,
            tone_score,
        )

    def test_create_session_builds_expected_workbook(self):
        workbook = load_workbook(self.workbook_path())
        self.assertEqual(workbook.sheetnames, [RESULT_SHEET, REQUEST_SHEET])
        self.assertEqual(
            [cell.value for cell in workbook[RESULT_SHEET][1]],
            RESULT_HEADERS,
        )
        self.assertEqual(workbook[REQUEST_SHEET].sheet_state, "hidden")
        self.assertEqual(
            self.service.get_current_session()["sessionId"],
            self.session["sessionId"],
        )

    def test_append_calculates_scores_and_writes_display_values(self):
        result = self.append()
        self.assertFalse(result["duplicate"])
        self.assertEqual(result["questionScore"], 70)
        self.assertEqual(result["totalScore"], 88)

        workbook = load_workbook(self.workbook_path(), data_only=False)
        values = [workbook[RESULT_SHEET].cell(2, column).value for column in range(1, 8)]
        self.assertEqual(
            values[:6],
            [
                "2026-09-21T18:30:00+08:00",
                "412345678",
                "7/8",
                "70/80",
                "18/20",
                "88/100",
            ],
        )
        self.assertTrue(values[6].startswith(f"transcripts/{self.session['sessionId']}/412345678_"))
        transcript_path = self.output_dir / Path(values[6])
        self.assertTrue(transcript_path.exists())
        self.assertEqual(workbook[RESULT_SHEET]["G2"].hyperlink.target, values[6])

    def test_transcript_uses_portable_relative_path_and_appends_text(self):
        result_id = str(uuid.uuid4())
        relative_path = self.service.prepare_student_transcript(
            self.session["sessionId"], result_id, "412345678"
        )
        self.assertFalse(Path(relative_path).is_absolute())
        self.assertIn("412345678_", Path(relative_path).name)

        self.service.append_student_transcript(
            self.session["sessionId"],
            result_id,
            "412345678",
            "2026-09-21T10:35:10+08:00",
            "step-3",
            "utterance-1",
            "不用怕，一下就好了",
        )
        transcript = (self.output_dir / relative_path).read_text(encoding="utf-8")
        self.assertIn("Student ID: 412345678", transcript)
        self.assertIn("step=step-3", transcript)
        self.assertIn("requestId=utterance-1", transcript)
        self.assertIn("不用怕，一下就好了", transcript)

    def test_student_id_is_stored_as_text_not_formula(self):
        self.append(student_id="=1+1")
        workbook = load_workbook(self.workbook_path(), data_only=False)
        student_cell = workbook[RESULT_SHEET]["B2"]
        self.assertEqual(student_cell.value, "=1+1")
        self.assertEqual(student_cell.data_type, "s")

    def test_same_result_id_is_idempotent(self):
        result_id = uuid.uuid4()
        first = self.append(result_id=result_id)
        second = self.append(result_id=result_id)
        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])

        workbook = load_workbook(self.workbook_path())
        self.assertEqual(workbook[RESULT_SHEET].max_row, 2)
        self.assertEqual(workbook[REQUEST_SHEET].max_row, 2)

    def test_same_result_id_with_different_payload_is_rejected(self):
        result_id = uuid.uuid4()
        self.append(result_id=result_id)
        with self.assertRaises(IdempotencyConflictError):
            self.append(result_id=result_id, tone_score=17)

    def test_new_session_makes_previous_session_inactive(self):
        previous_session_id = self.session["sessionId"]
        self.service.create_session(self.now)
        with self.assertRaises(InactiveSessionError):
            self.service.ensure_active_session(previous_session_id)

    def test_service_recovers_current_session_after_restart(self):
        restarted_service = ExcelResultService(self.output_dir)
        current = restarted_service.get_current_session()
        self.assertEqual(current["sessionId"], self.session["sessionId"])
        result = restarted_service.append_result(
            self.session["sessionId"],
            str(uuid.uuid4()),
            "400000001",
            "2026-09-21T18:45:00+08:00",
            8,
            20,
        )
        self.assertEqual(result["totalScore"], 100)

    def test_concurrent_submissions_do_not_lose_rows(self):
        def submit(index):
            return self.service.append_result(
                self.session["sessionId"],
                str(uuid.uuid4()),
                f"S{index:03d}",
                "2026-09-21T19:00:00+08:00",
                index % 9,
                index % 21,
            )

        with ThreadPoolExecutor(max_workers=6) as executor:
            results = list(executor.map(submit, range(12)))

        self.assertEqual(len(results), 12)
        workbook = load_workbook(self.workbook_path())
        self.assertEqual(workbook[RESULT_SHEET].max_row, 13)
        self.assertEqual(workbook[REQUEST_SHEET].max_row, 13)


class ResultModelTests(unittest.TestCase):
    def test_login_time_requires_timezone(self):
        with self.assertRaises(ValidationError):
            StudentResultSubmission(
                studentId="412345678",
                loginTime="2026-09-21T10:30:00",
                correctCount=7,
                toneScore=18,
            )

    def test_score_ranges_are_enforced(self):
        with self.assertRaises(ValidationError):
            StudentResultSubmission(
                studentId="412345678",
                loginTime="2026-09-21T10:30:00+08:00",
                correctCount=9,
                toneScore=18,
            )

    def test_boolean_scores_are_rejected(self):
        with self.assertRaises(ValidationError):
            StudentResultSubmission(
                studentId="412345678",
                loginTime="2026-09-21T10:30:00+08:00",
                correctCount=True,
                toneScore=18,
            )


if __name__ == "__main__":
    unittest.main()
